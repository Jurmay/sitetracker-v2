# -*- coding: utf-8 -*-
"""
Reads a priced BOQ spreadsheet into sections and items, ready to create
a project from a won tender.

Built for the BOQ_Pricing.xlsx files the tender dashboard writes for
every bid (Jurmay/bhutan-tender-dashboard, companies/<slug>/bids/<id>/):

    row 1   "BILL OF QUANTITIES: <title>"   (sometimes "FORM 7: ...")
    row 2   "NIT <number> · Bidder: <company> · ..."
    row 4   "GST registered? (Y/N)" | Y/N
    header  Sl. | Work | Description | Unit | Qty | Rate (Nu.) | Amount (Nu.) | notes
    lines   one per BOQ item, Sl. numbered; then Sub-total / GST / notes

Any other workbook with the same header row works too. Columns are found
by their header text, not their position. "Work" becomes the BOQ section:
"Work 1: Resurfacing of access road" and "Work 1" on later lines are the
same section (matched on the part before ":"; the longest name is kept).

The owner edits rates in Excel before submitting, so formulas may have
no saved result if the file was never opened in Excel; those lines (and
blank rates) are imported at rate 0 and listed as warnings.
"""
import io
import re

from openpyxl import load_workbook

HEADER_ALIASES = {
    "sl": ("sl", "sl.", "sl no", "s.no", "sno", "no", "no."),
    "work": ("work", "section", "boq"),
    "description": ("description", "description of items", "item description", "particulars"),
    "unit": ("unit", "units", "uom"),
    "quantity": ("qty", "quantity"),
    "rate": ("rate",),
}


# Summary rows under the table, which some files number like items.
SUMMARY_LINE = re.compile(r"^(sub[- ]?total|total|grand total|bhutan gst|gst\b)", re.I)


def _norm(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _header_key(text: str):
    t = _norm(text).lower()
    t = re.sub(r"\(.*?\)", "", t).strip()  # "Rate (Nu.)" -> "rate"
    for key, aliases in HEADER_ALIASES.items():
        if t in aliases:
            return key
    return None


def _number(value):
    """Number from a cell, or None (blank, text, or an unsaved formula)."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = _norm(value).replace(",", "")
    if text.startswith("="):
        return None
    try:
        return float(text)
    except ValueError:
        return None


class BoqImportError(ValueError):
    """The file isn't a BOQ we can read; the message says why, for the user."""


def parse_boq_xlsx(data: bytes) -> dict:
    try:
        values = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        raise BoqImportError("That file isn't an Excel workbook (.xlsx). Upload the BOQ_Pricing.xlsx from the tender dashboard.")

    ws = values[values.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]

    header_at, columns = None, {}
    for i, row in enumerate(rows[:40]):
        found = {}
        for j, cell in enumerate(row):
            key = _header_key(cell)
            if key and key not in found:
                found[key] = j
        if {"description", "unit", "quantity", "rate"} <= found.keys():
            header_at, columns = i, found
            break
    if header_at is None:
        raise BoqImportError(
            "Couldn't find the BOQ table. The first sheet needs a header row with "
            "Description, Unit, Qty and Rate columns (Sl. and Work are optional).")

    # Title, NIT, bidder and GST from the lines above the table.
    title = nit = bidder = None
    gst_registered = False
    for row in rows[:header_at]:
        text = " ".join(_norm(c) for c in row if c is not None)
        if not text:
            continue
        m = re.search(r"BILL OF QUANTITIES\s*:\s*(.+)", text, re.I)
        if m and not title:
            title = m.group(1).strip()
        m = re.search(r"\bNIT(?: No\.?)?\s+(.+?)\s*(?:·|$)", text)
        if m and not nit:
            nit = m.group(1).strip()
        m = re.search(r"Bidder:\s*(.+?)\s*(?:·|$)", text)
        if m and not bidder:
            bidder = m.group(1).strip()
        if re.search(r"GST registered", text, re.I):
            filled = [c for c in row if c is not None and _norm(c)]
            gst_registered = len(filled) > 1 and _norm(filled[-1]).upper().startswith("Y")

    sections, order, warnings = {}, [], []
    for row in rows[header_at + 1:]:
        cell = lambda key: row[columns[key]] if key in columns and columns[key] < len(row) else None
        sl = cell("sl")
        description = _norm(cell("description"))
        if not description or SUMMARY_LINE.match(description):
            continue  # blank lines, sub-totals, GST, grand total
        if "sl" in columns and _number(sl) is None:
            continue  # notes and headings without a line number
        if not _norm(cell("unit")) and cell("quantity") is None and cell("rate") is None:
            continue  # a heading row that happens to be numbered

        line_no = int(_number(sl)) if _number(sl) is not None else len(warnings) + 1
        unit = _norm(cell("unit")) or "LS"
        quantity = _number(cell("quantity"))
        rate = _number(cell("rate"))
        label = f"Line {line_no}: {description[:60]}{'…' if len(description) > 60 else ''}"
        if quantity is None:
            warnings.append(f"{label}: no quantity, imported as 0")
            quantity = 0.0
        if rate is None:
            warnings.append(f"{label}: no rate yet, imported as 0. Set the contract rate before work starts.")
            rate = 0.0

        work = _norm(cell("work")) or "General"
        key = work.split(":")[0].strip().lower() or "general"
        if key not in sections:
            sections[key] = {"name": work, "items": []}
            order.append(key)
        elif len(work) > len(sections[key]["name"]):
            sections[key]["name"] = work
        sections[key]["items"].append({
            "line_no": line_no,
            "description": description,
            "unit": unit,
            "quantity": round(quantity, 3),
            "rate": round(rate, 2),
            "amount": round(quantity * rate, 2),
        })

    result_sections = [sections[k] for k in order]
    if not any(s["items"] for s in result_sections):
        raise BoqImportError("The BOQ table has no item lines. Fill in the BOQ (the tender dashboard leaves it empty "
                             "when the tender's BOQ wasn't published) and upload it again.")

    total = round(sum(i["amount"] for s in result_sections for i in s["items"]), 2)
    return {
        "title": title,
        "tender_ref": nit,
        "bidder": bidder,
        "gst_registered": gst_registered,
        "sections": result_sections,
        "item_count": sum(len(s["items"]) for s in result_sections),
        "total": total,
        "warnings": warnings,
    }
